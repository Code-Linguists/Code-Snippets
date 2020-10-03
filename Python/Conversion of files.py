from docx import Document
from fpdf import FPDF
from pdfkit import from_file
import PyPDF2

class conversion:

    def text_to_pdf(text_file_path):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial Black",size = 14)
        x = open(text_file_path,"r")
        for l in x:
            pdf.cell(200,10,txt = l, ln = 1, aling = 'L')
        pdf.output('output.pdf')

    def text_to_docx(text_file_path):
        doc = Document()
        with open(text_file_path, 'r', encoding='utf-8') as file:
            doc.add_paragraph(file.read())
        doc.save("text.docx")

    def text_from_pdf(pdf_file_path):
        pdf_open = open(pdf_file_path, 'rb')
        pdf_read = PyPDF2.PdfFileReader(pdf_open)
        text = pdf_read.getPage(0)
        print(text.extractText())


